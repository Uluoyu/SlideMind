"""Bounded text extraction with source locations for supported document formats."""

from collections.abc import Iterator
from io import BytesIO
from typing import BinaryIO
from zipfile import ZipFile

from slidemind_parser.contracts import ParsedSection

MAX_INPUT_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 10_000
MAX_TEXT_CHARS = 2_000_000
SUPPORTED_TYPES = frozenset({"txt", "md", "pdf", "docx", "pptx"})


class _DocumentError(ValueError):
    """An actionable error that should pass through the parser boundary unchanged."""


def _check_text_size(size: int) -> None:
    if size > MAX_TEXT_CHARS:
        raise _DocumentError("提取文本超过大小限制，请拆分文档后重试。")


def _validate_archive(data: bytes) -> None:
    """Check declared and actual expansion before an OOXML library loads the ZIP."""
    with ZipFile(BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ARCHIVE_MEMBERS:
            raise _DocumentError("文档压缩包文件数量超过限制，请拆分文档后重试。")
        if sum(member.file_size for member in members) > MAX_ARCHIVE_BYTES:
            raise _DocumentError("文档解压大小超过限制，请拆分文档后重试。")
        expanded_bytes = 0
        for member in members:
            if member.flag_bits & 1:
                raise _DocumentError("文档已加密，请先移除密码后重新上传。")
            with archive.open(member) as contents:
                while block := contents.read(64 * 1024):
                    expanded_bytes += len(block)
                    if expanded_bytes > MAX_ARCHIVE_BYTES:
                        raise _DocumentError("文档解压大小超过限制，请拆分文档后重试。")


def _parse_pdf(data: bytes) -> Iterator[ParsedSection]:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as document:
        if document.needs_pass:
            raise _DocumentError("PDF 已加密，请先移除密码后重新上传。")
        for page_no, page in enumerate(document, start=1):
            yield ParsedSection(page.get_text("text", sort=True), page_no=page_no)


def _parse_docx(data: bytes) -> Iterator[ParsedSection]:
    from docx import Document
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(data))
    paragraph_index = 0
    table_index = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_index += 1
            yield ParsedSection(block.text, metadata={"paragraph_index": paragraph_index})
        else:
            table_index += 1
            for row_index, row in enumerate(block.rows, start=1):
                yield ParsedSection(
                    "\t".join(cell.text for cell in row.cells),
                    metadata={"table_index": table_index, "row_index": row_index},
                )


def _shape_text(shapes) -> Iterator[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_text(shape.shapes)
        elif shape.has_table:
            for row in shape.table.rows:
                yield "\t".join("" if cell.is_spanned else cell.text for cell in row.cells)
        elif shape.has_text_frame:
            yield shape.text_frame.text


def _parse_pptx(data: bytes) -> Iterator[ParsedSection]:
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))
    for slide_no, slide in enumerate(presentation.slides, start=1):
        texts = []
        text_size = 0
        for text in _shape_text(slide.shapes):
            if text.strip():
                text_size += len(text) + bool(texts)
                _check_text_size(text_size)
                texts.append(text)
        yield ParsedSection(
            "\n".join(texts), page_no=slide_no, metadata={"slide_no": slide_no}
        )


def parse_document(stream: BinaryIO, file_type: str) -> list[ParsedSection]:
    """Extract UTF-8 text or native PDF/Office text without closing the caller's stream.

    Page/slide and paragraph/table positions are one-based. DOCX has no reliable
    physical page numbers, so its sections use paragraph or table-row metadata.
    No OCR is performed; unsupported, damaged and textless inputs are rejected.
    """
    normalized_type = file_type.strip().lower().removeprefix(".")
    if normalized_type not in SUPPORTED_TYPES:
        raise _DocumentError(
            f"不支持的文档格式：{normalized_type}。请上传 TXT、MD、PDF、DOCX 或 PPTX 文件。"
        )
    try:
        data = stream.read(MAX_INPUT_BYTES + 1)
        if len(data) > MAX_INPUT_BYTES:
            raise _DocumentError("文件大小超过解析限制，请拆分文档后重试。")
        if not data:
            raise _DocumentError("文件为空，请上传包含文本的文档。")
        if normalized_type in {"txt", "md"}:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as error:
                raise _DocumentError("文本编码无效，请另存为 UTF-8 后重新上传。") from error
            if "\x00" in text:
                raise _DocumentError("文本格式无效，请另存为 UTF-8 纯文本后重新上传。")
            raw_sections = iter([ParsedSection(text)])
        elif normalized_type == "pdf":
            raw_sections = _parse_pdf(data)
        else:
            _validate_archive(data)
            raw_sections = _parse_docx(data) if normalized_type == "docx" else _parse_pptx(data)

        sections = []
        total_chars = 0
        for section in raw_sections:
            total_chars += len(section.text)
            _check_text_size(total_chars)
            text = section.text.replace("\r\n", "\n").replace("\r", "\n").strip()
            if text:
                sections.append(ParsedSection(
                    text,
                    page_no=section.page_no,
                    metadata={"source_type": normalized_type, **section.metadata},
                ))
        if not sections:
            raise _DocumentError(
                "文档没有可提取的文本；扫描件或纯图片文档请先进行 OCR 识别后重新上传。"
            )
        return sections
    except _DocumentError:
        raise
    except Exception as error:
        raise _DocumentError(
            f"{normalized_type.upper()} 文档损坏或格式不匹配，"
            "请确认文件能正常打开并重新导出后上传。"
        ) from error
