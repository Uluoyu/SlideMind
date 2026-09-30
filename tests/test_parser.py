from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from slidemind_parser.contracts import ParsedSection


def parse(data: bytes, file_type: str) -> list[ParsedSection]:
    from slidemind_parser.parsers import parse_document

    return parse_document(BytesIO(data), file_type)


def test_section_metadata_is_independent():
    first, second = ParsedSection("第一段"), ParsedSection("第二段")
    first.metadata["paragraph_index"] = 1
    assert second.metadata == {}


@pytest.mark.parametrize("file_type", ["txt", "md", ".TXT", " MD "])
def test_utf8_text_retains_chinese_and_markdown(file_type):
    sections = parse("\ufeff# 项目标题\r\n\r\n保留中文与 **格式**。\r\n".encode(), file_type)
    assert len(sections) == 1
    assert sections[0].text == "# 项目标题\n\n保留中文与 **格式**。"
    assert sections[0].page_no is None
    assert sections[0].metadata["source_type"] in {"txt", "md"}


def test_pdf_preserves_actual_page_numbers_including_empty_pages():
    import fitz

    with fitz.open() as document:
        document.new_page().insert_text((72, 72), "First page")
        document.new_page()
        document.new_page().insert_text((72, 72), "Third page")
        data = document.tobytes()
    sections = parse(data, "pdf")
    assert [section.page_no for section in sections] == [1, 3]
    assert [section.text for section in sections] == ["First page", "Third page"]


def test_docx_preserves_paragraph_table_order_and_source_positions():
    from docx import Document

    document = Document()
    document.add_paragraph("项目背景")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "产品"
    table.cell(0, 1).text = "营收"
    table.cell(1, 0).text = "引擎"
    table.cell(1, 1).text = "一百万"
    document.add_paragraph("项目结论")
    stream = BytesIO()
    document.save(stream)
    sections = parse(stream.getvalue(), "docx")
    assert [section.text for section in sections] == [
        "项目背景", "产品\t营收", "引擎\t一百万", "项目结论"
    ]
    assert all(section.page_no is None for section in sections)
    assert sections[0].metadata["paragraph_index"] == 1
    assert sections[1].metadata["table_index"] == 1
    assert sections[1].metadata["row_index"] == 1
    assert sections[2].metadata["row_index"] == 2
    assert sections[-1].metadata["paragraph_index"] == 2


def test_pptx_retains_slide_text_tables_and_slide_numbers():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    first.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = "第一张幻灯片"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    third = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = third.shapes.add_table(2, 2, 0, 0, Inches(4), Inches(2)).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "增长"
    table.cell(1, 1).text = "百分之十"
    stream = BytesIO()
    presentation.save(stream)
    sections = parse(stream.getvalue(), "pptx")
    assert [section.page_no for section in sections] == [1, 3]
    assert [section.metadata["slide_no"] for section in sections] == [1, 3]
    assert sections[0].text == "第一张幻灯片"
    assert sections[1].text == "指标\t结果\n增长\t百分之十"


@pytest.mark.parametrize("file_type", ["pdf", "docx", "pptx"])
def test_corrupt_document_has_actionable_chinese_error(file_type):
    with pytest.raises(ValueError, match="损坏|格式"):
        parse(b"this is not a document", file_type)


def test_xlsx_is_explicitly_unsupported():
    with pytest.raises(ValueError, match="不支持.*xlsx"):
        parse(b"", "xlsx")


@pytest.mark.parametrize("data", [b"", b" \r\n\t", "\ufeff".encode()])
def test_empty_text_has_actionable_chinese_error(data):
    with pytest.raises(ValueError, match="空|没有.*文本"):
        parse(data, "txt")


def test_non_utf8_text_suggests_resaving_as_utf8():
    with pytest.raises(ValueError, match="UTF-8"):
        parse("中文".encode("gbk"), "txt")


def test_image_only_pdf_suggests_ocr():
    import fitz

    with fitz.open() as document:
        page = document.new_page()
        pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
        pixmap.clear_with(128)
        page.insert_image(fitz.Rect(10, 10, 90, 90), pixmap=pixmap)
        data = document.tobytes()
    with pytest.raises(ValueError, match="OCR"):
        parse(data, "pdf")


@pytest.mark.parametrize("file_type", ["docx", "pptx"])
def test_empty_office_document_has_actionable_error(file_type):
    stream = BytesIO()
    if file_type == "docx":
        from docx import Document

        Document().save(stream)
    else:
        from pptx import Presentation

        Presentation().save(stream)
    with pytest.raises(ValueError, match="没有.*文本|空"):
        parse(stream.getvalue(), file_type)


def test_limits_extracted_text_instead_of_silently_truncating(monkeypatch):
    from slidemind_parser import parsers

    monkeypatch.setattr(parsers, "MAX_TEXT_CHARS", 10)
    with pytest.raises(ValueError, match="文本.*限制|文本.*过大"):
        parse("中文内容".encode() * 10, "txt")


def test_limits_compressed_office_expansion_before_parsing(monkeypatch):
    from slidemind_parser import parsers

    monkeypatch.setattr(parsers, "MAX_ARCHIVE_BYTES", 1000)
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "<document>" + "x" * 2000 + "</document>")
    with pytest.raises(ValueError, match="解压.*限制|解压.*过大"):
        parse(stream.getvalue(), "docx")


def test_limits_raw_input_bytes(monkeypatch):
    from slidemind_parser import parsers

    monkeypatch.setattr(parsers, "MAX_INPUT_BYTES", 8)
    with pytest.raises(ValueError, match="文件.*限制|文件.*过大"):
        parse(b"a" * 9, "txt")


def test_chunking_retains_page_position_and_exact_overlap():
    from slidemind_parser.chunking import chunk_sections

    source = ParsedSection("甲乙丙丁戊己庚辛壬癸子丑", page_no=3, metadata={"slide_no": 3})
    chunks = chunk_sections([source], max_chars=6, overlap=2)
    assert [chunk.text for chunk in chunks] == ["甲乙丙丁戊己", "戊己庚辛壬癸", "壬癸子丑"]
    assert [chunk.page_no for chunk in chunks] == [3, 3, 3]
    assert [chunk.metadata["slide_no"] for chunk in chunks] == [3, 3, 3]
    assert [chunk.metadata["char_start"] for chunk in chunks] == [0, 4, 8]
    assert [chunk.metadata["char_end"] for chunk in chunks] == [6, 10, 12]
    assert [chunk.metadata["chunk_index"] for chunk in chunks] == [0, 1, 2]
    assert all(chunk.metadata["section_index"] == 0 for chunk in chunks)
    assert source.metadata == {"slide_no": 3}
    assert chunk_sections([source], max_chars=6, overlap=2) == chunks


def test_chunking_long_chinese_stays_bounded_and_does_not_cross_source_pages():
    from slidemind_parser.chunking import chunk_sections

    sections = [ParsedSection("中" * 1900, page_no=1), ParsedSection("另一页", page_no=2)]
    chunks = chunk_sections(sections)
    assert [len(chunk.text) for chunk in chunks] == [900, 900, 300, 3]
    assert [chunk.page_no for chunk in chunks] == [1, 1, 1, 2]
    assert chunks[-1].metadata["section_index"] == 1
    assert chunks[-1].metadata["char_start"] == 0
    assert chunks[-1].metadata["char_end"] == 3


def test_chunking_without_overlap_keeps_whitespace_and_skips_blank_sections():
    from slidemind_parser.chunking import chunk_sections

    sections = [ParsedSection(" \n"), ParsedSection("甲 乙\n丙")]
    chunks = chunk_sections(sections, max_chars=3, overlap=0)
    assert [chunk.text for chunk in chunks] == ["甲 乙", "\n丙"]
    assert all(chunk.metadata["section_index"] == 1 for chunk in chunks)
    assert chunk_sections([]) == []


@pytest.mark.parametrize("max_chars,overlap", [(0, 0), (-1, 0), (5, -1), (5, 5), (5, 6)])
def test_chunking_rejects_invalid_window_parameters(max_chars, overlap):
    from slidemind_parser.chunking import chunk_sections

    with pytest.raises(ValueError):
        chunk_sections([ParsedSection("中文")], max_chars=max_chars, overlap=overlap)
